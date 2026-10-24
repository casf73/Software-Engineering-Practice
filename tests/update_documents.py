"""同步三轮需求、测试 Word/Markdown、接口及数据模型说明。"""
from pathlib import Path
from datetime import datetime
from collections import Counter
from urllib.parse import quote
import json,re,xml.etree.ElementTree as ET
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import importlib.util
ROOT=Path(__file__).resolve().parent.parent
TESTS=ROOT/'tests';SOURCE=ROOT/'project/mvp-v1.0';DATE='2026-10-24'
spec=importlib.util.spec_from_file_location('fixtures',TESTS/'generate-iteration-tests.py');gen=importlib.util.module_from_spec(spec);spec.loader.exec_module(gen)
RAW='https://raw.githubusercontent.com/casf73/Software-Engineering-Practice/main/'
ROUND_NAMES={1:'首轮',2:'第二轮',3:'第三轮'}
TITLES={1:'在线购物系统首轮验收测试用例',2:'在线购物系统第二轮验收测试用例',3:'在线购物系统第三轮验收测试用例'}
def save(path,text):
 path.parent.mkdir(parents=True,exist_ok=True);path.write_text(text.strip()+'\n',encoding='utf-8')
def table(head,rows):return '\n'.join(['| '+' | '.join(head)+' |','| '+' | '.join(['---']*len(head))+' |']+['| '+' | '.join(str(v).replace('\n','<br>').replace('|','\\|') for v in row)+' |' for row in rows])
def compact(value):
 if isinstance(value,str) and len(value)>65 and len(set(value))<=3:return f'{value[0]}重复{len(value)}次'
 if isinstance(value,list) and len(value)>3 and all(x==value[0] for x in value):return f'{compact(value[0])}共{len(value)}项'
 if isinstance(value,dict):return '；'.join(f'{k}={compact(v)}' for k,v in value.items())
 if isinstance(value,list):return '；'.join(compact(x) for x in value) if value else '[]'
 return json.dumps(value,ensure_ascii=False) if value is None or isinstance(value,bool) else str(value)
def result_for(r,obj):
 if obj['mode']=='manual':return '待人工执行'
 path=SOURCE/'backend/target/acceptance-results'/str(r)/(obj['id']+'.json')
 if path.exists():return json.loads(path.read_text(encoding='utf-8')).get('result','待执行')
 if r==1 and int(obj['id'].split('-')[-1])<=3:
  p=SOURCE/'backend/target/surefire-reports/TEST-cn.codelyy.store.StoreFlowIntegrationTest.xml'
  if p.exists():
   tree=ET.parse(p).getroot();case=obj['caseId'];idx=int(obj['id'][-2:]);methods={'T01':'sellerLoginAndAuthorization','T02':'productCreationAndValidation','T03':'queuePasscodesAndSellerActions','T04':'tradeFreezeAndCancellationBoundaries','T05':'tradeFailureVoidsIntent','T06':'tradeFailureRequeuesIntent','T07':'successfulTradeEndsProduct','T08':'sellerDelistsProduct','T09':'imageUploadsAndLimits','T10':'legacyMultipleImagesReturnFirst'}
   # 原始参数化实例全部通过时才把对应30项登记为通过。
   if tree.get('failures')=='0' and tree.get('errors')=='0' and int(tree.get('tests','0'))>=31:return '通过'
 return '待执行'
def manifest():
 data=[]
 for r,base in gen.ROUNDS.items():
  records=[]
  for p in sorted(base.rglob('instance.json')):
   o=json.loads(p.read_text(encoding='utf-8'));o['file']=p.relative_to(TESTS).as_posix();o['result']=result_for(r,o);inputs=[];expected=[];steps=[];requests=[]
   if o['mode']=='manual':
    manual_data=json.loads((p.parent/o['dataFile']).read_text(encoding='utf-8'));
    if '输入集合' in manual_data:inputs=[compact(x['输入']) for x in manual_data['输入集合'] if not (x['输入'].get('username')=='admin' and x['输入'].get('password')=='123456')]
    else:inputs=[compact(manual_data)]
    if manual_data.get('本实例素材文件'):inputs.append('本实例素材：'+','.join(manual_data['本实例素材文件']))
    steps=o['steps'];expected=[o['expected']]
   elif o.get('steps'):
    for i,st in enumerate(o['steps'],1):
     body=json.loads((p.parent/st['bodyFile']).read_text(encoding='utf-8')) if 'bodyFile' in st else {}
     if body and not (st['path']=='/api/seller/login' and body.get('password')=='123456'):inputs.append(compact(body))
     if st.get('file'):inputs.append(f'文件={st["file"]}，MIME={st.get("mime")}，字节={ (p.parent/st["file"]).stat().st_size }')
     steps.append(f'{i}. {st["method"]} {st["path"]} => {st["status"]}')
     if st.get('assertions'):expected.append(compact(st['assertions']))
     if st.get('absent'):expected.append('不返回'+','.join(st['absent']))
     requests.append([o['id'],i,st['method'],st['path'],st.get('session','访客'),json.dumps(body,ensure_ascii=False),st['status'],json.dumps(st.get('assertions',{}),ensure_ascii=False),st.get('bodyFile',st.get('file',''))])
   else:
    inputs=[compact(o.get('input',{}))];steps=['读取该实例 input 字段，执行 StoreFlowIntegrationTest 中对应编号的参数化方法。'];expected=[o['expected']]
   o['dataSummary']='；'.join(dict.fromkeys(inputs)) or '无请求体；参数见操作步骤'
   o['expectedSummary']='；'.join(dict.fromkeys(expected)) or o['expected']
   o['stepSummary']='\n'.join(steps);o['requests']=requests;records.append(o)
   # 保留原始人工通过记录已转为历史；新整理版本逐项登记本次执行状态。
   stored={k:v for k,v in o.items() if k not in ('file','dataSummary','expectedSummary','stepSummary','requests')};gen.write(p,stored)
  data.append({'round':r,'title':TITLES[r],'records':records,'cases':gen.CASES[r]})
 return data
FIGURES=[('业务流程图-商品与交易主流程.png','商品与交易主流程',['卖家发布多个商品，各自独立库存','买家选择商品，访客或客户提交意向','每件商品独立 FIFO 排队','卖家开始队首交易，暂停该商品新意向','成功扣减1件；库存为0时下架','失败作废或重排；有库存继续在售']),('业务流程图-买家浏览、排队与口令码操作.png','买家浏览与意向操作',['浏览、关键词搜索、两级分类筛选','访客填写姓名电话；客户填写交易地点','提交到指定商品队列并领取口令码','客户记录保存到自己的购买历史','等待者可修改姓名电话或撤销','交易者不可修改撤销；结束后口令失效']),('业务流程图-卖家后台与历史记录.png','卖家后台与客户历史',['卖家登录后台','管理多个商品、库存和两级分类','处理每件商品队列与交易结果','客户管理搜索用户名或电话','查看注册客户及其购买历史','展示交易地点和快照，不返回密码哈希']),('业务流程图-商品发布与媒体处理.png','商品发布与媒体处理',['填写名称价格库存及分类','相册选择0至9张JPG或PNG','每张图片不超过5MB','编辑富文本与本地描述图片','服务端清理危险标签及图片地址','保存后公开展示，库存为0不在售']),('状态图-商品状态.png','商品状态',['库存大于0发布为ON_SALE','手动冻结进入MANUALLY_FROZEN','手动解冻返回ON_SALE','开始交易进入IN_TRADE','成功后有库存或失败：返回ON_SALE','售罄或主动下架：DELISTED，可补货上架']),('状态图-购买意向状态.png','购买意向状态',['提交意向进入WAITING','队首开始交易进入IN_TRADE','成功进入TRADE_SUCCESS','失败作废进入TRADE_FAILED，或重排WAITING','等待者撤销CANCELLED；卖家作废VOIDED','售罄结束剩余等待；主动下架PRODUCT_DELISTED']),('状态图-口令码有效性.png','口令码有效性',['提交意向生成唯一口令码','等待、冻结、交易中保持有效','修改姓名电话保留原码及排位','失败重排保留原码，移动到队尾','撤销、作废、成功、售罄或下架使相应码失效','失效不删除客户购买历史和商品快照'])]
def diagrams():
 font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',25);small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',21)
 for filename,title,labels in FIGURES:
  im=Image.new('RGB',(1200,680),'white');d=ImageDraw.Draw(im);d.text((50,28),title+'  第三轮',font=font,fill='#111827')
  for i,text in enumerate(labels):
   y=90+i*90;d.rounded_rectangle((65,y,1135,y+65),radius=8,outline='#aeb9c7',width=2,fill='#f8fafc');d.text((88,y+17),text,font=small,fill='#222f40')
   if i<5:d.line((600,y+65,600,y+86),fill='#6b7280',width=2);d.polygon([(594,y+80),(606,y+80),(600,y+88)],fill='#6b7280')
  im.save(ROOT/'process-and-state-diagrams'/filename)
def docx_from_md(path,md,landscape=False):
 doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.65);sec.right_margin=Inches(.65)
 if landscape:sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
 else:sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
 for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.size=Pt(10 if name=='Normal' else 21 if name=='Title' else 14 if name=='Heading 1' else 12);st.font.color.rgb=RGBColor(0,0,0);st.element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
  st.paragraph_format.space_after=Pt(6)
 for st in doc.styles:
  if st.type==1 and st.element.pPr is not None:
   for border in list(st.element.pPr.findall(qn('w:pBdr'))):st.element.pPr.remove(border)
 lines=md.splitlines();i=0;code=False
 while i<len(lines):
  line=lines[i].strip()
  if line.startswith('```'):code=not code;i+=1;continue
  if not line:i+=1;continue
  if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:|-]+\|$',lines[i+1].strip()):
   raw=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    l=lines[i].strip()
    if not re.match(r'^\|[\s:|-]+\|$',l):raw.append([x.strip().replace('<br>','\n').replace('`','') for x in re.split(r'(?<!\\)\|',l.strip('|'))])
    i+=1
   cols=len(raw[0]);t=doc.add_table(rows=0,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;t.style='Table Grid'
   borders=OxmlElement('w:tblBorders')
   for edge in ['top','left','bottom','right','insideH','insideV']:
    border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
   t._tbl.tblPr.append(borders)
   usable=sec.page_width-sec.left_margin-sec.right_margin
   ratios=([.16,.50,.22,.12] if cols==4 and raw[0][0]=='实例与场景' else [.12,.40,.24,.24] if cols==4 else [.12,.18,.17,.25,.28] if cols==5 else [.2,.8] if cols==2 else [1/cols]*cols)
   for rown,values in enumerate(raw):
    if rown and raw[0][0]=='实例与场景':values[-1]=re.sub(r'(?:验收测试数据|第二轮验收测试数据|第三轮验收测试数据)/[^/]+/[^/]+/', '', values[-1])
    row=t.add_row();pr=row._tr.get_or_add_trPr()
    if max(map(len,values),default=0)<600:pr.append(OxmlElement('w:cantSplit'))
    if rown==0:pr.append(OxmlElement('w:tblHeader'))
    for j,cell in enumerate(row.cells):
     cell.width=int(usable*ratios[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=values[j] if j<len(values) else ''
     tcpr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
     for edge in ['top','left','bottom','right']:
      z=OxmlElement('w:'+edge);z.set(qn('w:w'),'85');z.set(qn('w:type'),'dxa');m.append(z)
     tcpr.append(m)
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF5' if rown==0 else 'FFFFFF');tcpr.append(shade)
     for p in cell.paragraphs:
      p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=False
      for run in p.runs:run.font.size=Pt(9);run.bold=rown==0
   for j,col in enumerate(t.columns):col.width=int(usable*ratios[j])
   doc.add_paragraph();continue
  image=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  if image:
   filename=image.group(2).split('/')[-1];from urllib.parse import unquote
   local=ROOT/'process-and-state-diagrams'/unquote(filename)
   if local.exists():doc.add_picture(str(local),width=Inches(6.65 if not landscape else 9.8));doc.add_paragraph(image.group(1),'Caption')
   i+=1;continue
  line=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',line).replace('`','')
  if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
  elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
  elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
  else:doc.add_paragraph(line)
  i+=1
 footer=sec.footer.paragraphs[0];footer.text='在线购物系统  '+DATE+'  第';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run('页')
 doc.save(path)
def requirement():
 terms=[['访客','不注册也可浏览商品、按姓名电话提交意向，通过口令码管理。'],['注册客户','填写用户名、密码、电话及默认交易地点注册，登录后可保存和查看自己的购买记录。'],['卖家','唯一后台账号，管理商品、分类、库存、交易和注册客户。'],['选中商品','列表中正在查看或操作的商品，可同时存在多个有效商品。'],['购买意向','针对一个指定商品的一件购买申请；每件商品独立FIFO队列。'],['FIFO','按服务端queue_at升序，相同时间按意向编号升序。'],['手动冻结','只暂停该商品接收新意向，保留队列，可由卖家解除。'],['交易冻结','该商品已有队首交易，必须先登记结果，不能手动解除。'],['库存','0～1000000整数；成功减1，失败不减，归零自动下架。'],['口令码','单条意向的持有者凭据；意向结束后置空，历史不删除。'],['商品快照','客户下单时保存的商品名称和价格，后续改商品不回写快照。'],['两级分类','一级与二级父子分类，禁止三级、循环、同级重名。'],['富媒体描述','允许的格式化文字及本地描述图片；危险标签及图片地址清理或拒绝。']]
 text=f'''# 在线购物系统需求规格说明书

版本 V3.0；修订日期 {DATE}。本说明整合三轮功能，第三轮验收挑战需求 A.1 与 A.2，已实现规则以本版为准。

## 术语表

{table(['术语','定义'],terms)}

## 1 迭代范围

{table(['轮次','验收内容'],[['第一次','访客浏览、单件商品基本发布、姓名电话与口令码、FIFO、交易成功失败、主动下架、图片基础上传。'],['第二次','等待买家修改姓名电话；手动冻结和解冻；卖家修改密码及旧会话撤销。'],['第三次','A.1 多商品、两级分类、库存、搜索、多图富媒体；A.2 客户注册登录、客户历史、卖家客户查询。']])}

## 2 角色与权限

访客可免登录浏览、搜索、筛选商品并提交意向；注册客户保留访客能力，登录后可查看自己账号提交的购买记录。卖家必须登录才能管理商品、队列、分类、客户和客户历史。

客户历史的归属只使用服务端会话中的客户编号，不接受客户端指定其他客户编号；同名或同电话访客意向不能自动认领。卖家客户查询不返回密码、密码哈希或客户口令码。

## 3 商品和分类需求

{table(['字段或能力','规则'],[['同时发布','允许多个商品同时在售；冻结、交易、库存与等待队列按商品隔离。'],['名称','必填，最多80字符。'],['价格','大于0，最多2位小数，整数部分最多10位。'],['库存','0～1000000整数；0库存商品发布即下架；交易中禁止编辑。'],['分类','商品可关联一个一级或二级分类，也可未分类；一级筛选包含子类。'],['分类维护','新增、编辑、删除；禁止三级、循环与同级重名，已被商品或子类引用时拒绝删除。'],['图片','0～9张JPG/PNG，每张不超过5MB；保持相册顺序，前端可切换缩略图。'],['描述','最多20000字符，清理后仍检查长度；支持标题、段落、加粗、斜体、列表及本地图片。'],['浏览和搜索','按名称或描述可见文字搜索，支持分类筛选、分页；每页1～100条，页码从0开始。']])}

第三轮替代首轮的“同一时间唯一商品”“最多一张图片”和“每次成功必然下架”规则。多标签、商品视频和视频自动删除仍属于后续需求，不计入三轮已实现功能。

## 4 购买意向与队列

每条意向对应1件指定商品，姓名必填且最多80字符；电话须含数字，允许数字、加号、括号、空格、短横线，总长度6～30字符。仅在售且有库存的商品接受新意向。

访客提交成功领取唯一口令码，注册客户提交还必须确认交易地点，长度1～200字符，默认填入注册地点；本次地点修改不修改账号默认地点。服务端保存customer_id、product_name_snapshot、price_snapshot和meeting_location。

队列按每个商品的queue_at和id排序，卖家只能开始队首交易。等待中凭码可修改姓名电话或撤销；修改保留编号、提交时间、原码及排位。交易中的买家不能修改或撤销；同一商品其他等待者仍可修改或撤销。

## 5 交易和状态

{table(['操作','结果'],[['开始交易','商品ON_SALE转IN_TRADE；队首WAITING转IN_TRADE，拒绝该商品新意向。'],['成功且剩余库存大于0','当前意向TRADE_SUCCESS并清空口令码，库存减1，商品恢复ON_SALE；其他等待者继续排队。'],['成功后库存为0','商品DELISTED，当前意向成功；其他等待意向结束为TRADE_FAILED并清空口令码。'],['失败并作废','当前意向TRADE_FAILED，口令码失效，库存不变，商品恢复ON_SALE。'],['失败并重排','当前意向回WAITING，保留原码并更新queue_at到队尾；库存不变。'],['手动冻结与解冻','ON_SALE与MANUALLY_FROZEN之间转换；原等待队列和口令码保留。'],['主动下架','只允许ON_SALE；等待者转PRODUCT_DELISTED，口令码失效。'],['补货上架','DELISTED且库存大于0允许重新上架；历史意向不自动恢复。']])}

卖家作废等待者使用VOIDED状态；买家撤销使用CANCELLED状态。没有超时自动成交或自动处理下一位买家。已结束意向和客户历史持久保留。

## 6 客户升级 A.2

{table(['编号','验收要求'],[['A.2-1','客户可自由注册；用户名3～32位Unicode文字、数字、下划线或短横线，大小写不区分重复；密码8～128位，电话符合规则，默认交易地点必填且最多200字符。注册成功自动登录。'],['A.2-2','客户登录后查看本人购买历史，包含等待、交易、成功、失败、撤销、作废和下架结束记录；按提交时间和编号倒序分页。显示商品快照、联系人、交易地点、状态、时间及有效排位。'],['A.2-3','卖家查看所有注册客户的用户名、电话、默认地点、注册时间；可按用户名或电话搜索，并分页查看指定客户购买历史。']])}

## 7 登录和密码

卖家初始账号admin/123456仅用于首次建库。卖家改密验证原密码，新密码8～128位，增加会话版本使所有旧卖家会话失效。

卖家和客户密码均使用PBKDF2-HMAC-SHA256、600000次迭代、16字节独立随机盐、32字节派生结果存储。HTTP Session保存角色身份；登录更换会话编号；客户与卖家身份分开管理，客户退出只清除客户身份，卖家退出或改密只清除卖家身份。

## 8 数据与实现边界

数据库为H2文件库，日常运行保留data目录。JPA根据实体更新结构，新字段可空以兼容旧访客意向。产品、分类、卖家账号、客户账号、购买意向共五类实体。

客户购买历史使用当前意向记录及下单快照，不代表每次重排和每次交易尝试均有独立事件表。完整操作审计、商品视频、多标签、客户自助修改默认地点等未实现部分不列为第三轮通过项。

## 9 业务流程与状态图

下列七张图按第三轮业务规则更新，为各流程和状态规则示意。

'''
 for name,title,_ in FIGURES:text+=f'### {title}\n\n![{title}]({RAW}process-and-state-diagrams/{quote(name)})\n\n'
 text+='## 10 验收与测试\n\n三轮测试均包含总方案、总用例、每总用例至少8个详细实例。功能测试和接口测试分目录，接口使用JUnit、MockMvc及独立内存H2读取实例文件；页面功能测试另记人工结果。边界值、等价类、状态转换和权限隔离等方法见各轮测试文档及Excel末行。\n'
 save(ROOT/'requirement/在线购物系统需求规格说明书.md',text)
 docx_from_md(ROOT/'requirement/在线购物系统需求规格说明书.docx',text)
 third='''# 第三轮需求补充说明

版本 V3.0，日期 2026-10-24。本轮交付挑战需求A.1和A.2；第二轮三项基础功能继续保留。原第二轮资料中的A.1实现素材统一归入本轮验收。

'''+text[text.index('## 3 商品和分类需求'):text.index('## 9 业务流程与状态图')]
 save(ROOT/'requirement/第三轮/第三轮需求补充说明.md',third)
 p=ROOT/'requirement/第二轮/第二轮需求补充说明.md'
 save(p,'''# 第二轮需求补充说明

版本 V2.3，整理日期2026-10-24。第二轮仅验收以下三项基础改进；挑战需求A.1和A.2在第三轮交付，不回填为第二轮完成项。

'''+table(['编号','需求'],[['R201','等待买家凭口令码修改姓名和电话，保留原编号、口令码和排位；手动冻结和他人交易期间仍可改，当前交易者不能改。'],['R202','卖家手动冻结保留等待队列、拒绝新意向；手动解冻恢复在售，不能手动解除交易冻结。'],['R203','改密校验当前密码，新密码8～128位；PBKDF2加盐持久化，修改后旧卖家会话失效。']])+'\n\n对应测试包含功能测试和接口测试，每总用例各8个实例。')
 return text
ENDPOINTS=[
('GET','/api/public/products','公开','按关键词、分类、页码浏览','PageView<ProductView>'),('GET','/api/public/products/{id}','公开','指定商品详情','ProductView'),('POST','/api/public/products/{id}/intents','公开','访客提交指定商品意向','IssuedIntentView'),('GET','/api/public/categories','公开','两级分类列表','CategoryView[]'),('GET','/api/public/products/current','公开','兼容唯一当前商品','ProductView'),('POST','/api/public/intents','公开','兼容当前商品意向提交','IssuedIntentView'),('GET','/api/public/intents/{passcode}','口令码','意向状态与排位','BuyerIntentView'),('PUT','/api/public/intents/{passcode}','口令码','修改等待者联系信息','BuyerIntentView'),('DELETE','/api/public/intents/{passcode}','口令码','撤销意向','204 无正文'),
('POST','/api/seller/login','公开','卖家登录','{authenticated:boolean}'),('GET','/api/seller/session','公开','卖家会话状态','{authenticated:boolean}'),('POST','/api/seller/logout','卖家','退出卖家身份','204 无正文'),('PUT','/api/seller/password','卖家','改密并撤销旧卖家会话','204 无正文'),('GET','/api/seller/products','卖家','包含下架商品的分页列表','PageView<ProductView>'),('POST','/api/seller/products','卖家','发布商品','ProductView'),('PUT','/api/seller/products/{id}','卖家','编辑指定商品','ProductView'),('GET','/api/seller/products/{id}/workbench','卖家','商品、队列及当前交易','WorkbenchView'),('POST','/api/seller/products/{id}/freeze','卖家','手动冻结','ProductView'),('POST','/api/seller/products/{id}/unfreeze','卖家','手动解冻','ProductView'),('POST','/api/seller/products/{id}/relist','卖家','补货后重新上架','ProductView'),('POST','/api/seller/products/{id}/delist','卖家','主动下架','204 无正文'),('POST','/api/seller/products/{id}/trades/start','卖家','开始队首交易','IntentView'),('POST','/api/seller/products/image','卖家','multipart图片上传','{url:string}'),('POST','/api/seller/categories','卖家','创建分类','CategoryView'),('PUT','/api/seller/categories/{id}','卖家','编辑分类','CategoryView'),('DELETE','/api/seller/categories/{id}','卖家','删除无引用分类','204 无正文'),('POST','/api/seller/trades/{intentId}/result','卖家','登记交易结果','204 无正文'),('POST','/api/seller/intents/{intentId}/action','卖家','作废或重新排队','204 无正文'),('GET','/api/seller/workbench','卖家','兼容当前商品工作台','WorkbenchView'),('PUT','/api/seller/products/current','卖家','兼容当前商品编辑','ProductView'),('POST','/api/seller/products/current/freeze','卖家','兼容当前商品冻结','ProductView'),('POST','/api/seller/products/current/unfreeze','卖家','兼容当前商品解冻','ProductView'),('POST','/api/seller/products/current/delist','卖家','兼容当前商品下架','204 无正文'),('POST','/api/seller/trades/start','卖家','兼容当前商品队首交易','IntentView'),
('POST','/api/customer/register','公开','注册客户并自动登录','CustomerSessionView'),('POST','/api/customer/login','公开','客户登录','CustomerSessionView'),('GET','/api/customer/session','公开','查询客户会话','CustomerSessionView'),('POST','/api/customer/logout','客户','退出客户身份','204 无正文'),('GET','/api/customer/orders','客户','只查看本人历史','PageView<CustomerOrderView>'),('POST','/api/customer/products/{id}/intents','客户','保存本次地点和客户归属','IssuedIntentView'),('GET','/api/seller/customers','卖家','搜索注册客户','PageView<CustomerView>'),('GET','/api/seller/customers/{id}','卖家','指定客户详情','CustomerView'),('GET','/api/seller/customers/{id}/orders','卖家','指定客户历史','PageView<CustomerOrderView>')]
ENTITIES={'categories':['id BIGINT PK','name VARCHAR(80)','parent_id BIGINT FK nullable'],'products':['id BIGINT PK','name VARCHAR(80)','price DECIMAL(12,2)','description VARCHAR(20000)','image_urls VARCHAR(6000)','stock INTEGER','category_id BIGINT FK','status VARCHAR(20)','active_slot BOOLEAN legacy','created_at / updated_at TIMESTAMP'],'purchase_intents':['id BIGINT PK','product_id BIGINT FK','customer_id BIGINT FK nullable','product_name_snapshot VARCHAR(80)','price_snapshot DECIMAL(12,2)','meeting_location VARCHAR(200)','name VARCHAR(80)','phone VARCHAR(30)','passcode VARCHAR(24) UNIQUE nullable','status VARCHAR(24)','submitted_at / queue_at TIMESTAMP'],'customer_accounts':['id BIGINT PK','username VARCHAR(32)','username_key VARCHAR(32) UNIQUE','password_hash VARCHAR(256)','phone VARCHAR(30)','default_location VARCHAR(200)','created_at TIMESTAMP'],'seller_accounts':['id BIGINT PK = 1','password_hash VARCHAR(256)','session_version BIGINT']}
def api_and_model():
 model='# 后端数据模型图\n\n版本V3.0；日期'+DATE+'。第三轮新增客户账号及客户意向归属、地点和商品快照字段。\n\n![数据模型图]('+RAW+'project/mvp-v1.0/backend/docs/'+quote('数据模型图.png')+')\n\n## 实体关系\n\n```mermaid\nerDiagram\n    categories o|--o{ categories : parent_id\n    categories o|--o{ products : category_id\n    products ||--o{ purchase_intents : product_id\n    customer_accounts o|--o{ purchase_intents : customer_id\n'
 for entity,fields in ENTITIES.items():
  model+=f'    {entity} {{\n'
  for field in fields:
   for name in field.split(' / ') if ' / ' in field else [field]:
    parts=name.split();fname=parts[0].replace(',','');typ='DECIMAL' if 'DECIMAL' in field else 'TIMESTAMP' if 'TIMESTAMP' in field else 'BIGINT' if 'BIGINT' in field else 'INTEGER' if 'INTEGER' in field else 'VARCHAR';model+=f'        {typ} {fname}\n'
  model+='    }\n'
 model+='```\n\n## 第三轮新增结构\n\n'+table(['表或字段','含义'],[['customer_accounts','新增客户账号表，大小写归一用户名唯一，密码只存PBKDF2哈希。'],['purchase_intents.customer_id','可空客户外键；既有访客意向保持空，不凭姓名电话回填。'],['product_name_snapshot / price_snapshot','下单时商品名称与价格快照，后续商品编辑不修改。'],['meeting_location','本次交易地点，注册客户提交必填；默认来自客户账号。'],['idx_intent_customer_submitted','customer_id、submitted_at、id索引，支持客户历史分页。']])+'\n\n卖家账号为固定系统身份，与客户账号分表；密码哈希不放入任何接口DTO。\n'
 save(SOURCE/'backend/docs/数据模型图.md',model)
 # 以实际字段绘制可阅读的模型图，关系文字置于图底。
 font='C:/Windows/Fonts/msyh.ttc';im=Image.new('RGB',(1600,1050),'white');d=ImageDraw.Draw(im);titlefont=ImageFont.truetype(font,27);fieldfont=ImageFont.truetype(font,20)
 cards=[(30,60,'categories'),(560,60,'products'),(1090,60,'purchase_intents'),(30,640,'customer_accounts'),(1090,740,'seller_accounts')];svg=['<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="1050"><rect width="1600" height="1050" fill="white"/>']
 for x,y,name in cards:
  fields=ENTITIES[name];height=65+len(fields)*31;d.rectangle((x,y,x+480,y+height),outline='#94a3b8',width=2);d.rectangle((x,y,x+480,y+43),fill='#e8eef5');d.text((x+14,y+7),name,font=titlefont,fill='#111827')
  svg.append(f'<rect x="{x}" y="{y}" width="480" height="{height}" fill="#f8fafc" stroke="#94a3b8"/><text x="{x+14}" y="{y+30}" font-size="25" font-family="Microsoft YaHei">{name}</text>')
  for i,field in enumerate(fields):d.text((x+14,y+55+i*31),field,font=fieldfont,fill='#182536');svg.append(f'<text x="{x+14}" y="{y+76+i*31}" font-size="18" font-family="Microsoft YaHei">{field.replace("&","&amp;")}</text>')
 d.text((550,650),'categories 1 → N products',font=fieldfont,fill='#182536');d.text((550,690),'products 1 → N purchase_intents',font=fieldfont,fill='#182536');d.text((550,730),'customer_accounts 1 → N purchase_intents',font=fieldfont,fill='#182536');d.text((550,770),'customer_id 可空，访客历史不自动绑定',font=fieldfont,fill='#182536')
 svg.append('<text x="550" y="850" font-size="21" font-family="Microsoft YaHei">categories 1:N products；products 1:N purchase_intents</text><text x="550" y="890" font-size="21" font-family="Microsoft YaHei">customer_accounts 1:N purchase_intents；访客 customer_id 为空</text></svg>')
 im.save(SOURCE/'backend/docs/数据模型图.png');save(SOURCE/'backend/docs/数据模型图.svg',''.join(svg))
 text=f'# 在线购物系统后端接口说明\n\n版本V3.0；日期{DATE}。当前生产源码共{len(ENDPOINTS)}条业务路由（含兼容接口）；第三轮新增客户注册登录、本人购买历史、卖家客户查询。\n\n## 数据模型图\n\n![数据模型图]({RAW}project/mvp-v1.0/backend/docs/{quote("数据模型图.png")})\n\n详细结构见[数据模型说明](docs/数据模型图.md)。\n\n## 接口总览\n\n'+table(['方法','接口','身份','用途','返回类型'],ENDPOINTS)
 text+='''\n\n## 通用约定

基址为http://localhost:8080，JSON请求使用Content-Type: application/json，图片上传为multipart/form-data的file字段。金额JSON为number，编号为number，状态为string，日期为ISO本地日期时间字符串；可空字段返回null。分页请求page从0开始，size为1～100，返回items、total、page、size。

成功查询或更新通常200，创建201，删除/退出/登记结果204无正文；400参数校验失败，401身份无效，404资源不存在或口令失效，409重复账号或业务状态冲突。卖家和客户使用同源HTTP Session Cookie，按独立身份字段授权。

## 第三轮请求模型

'''+table(['模型','请求字段及约束'],[['ProductRequest','name必填≤80；price>0且最多2位小数；description≤20000；images≤9；stock整数0～1000000；categoryId可空。'],['CategoryRequest','name必填≤80；parentId可空，只能指向一级分类。'],['CustomerRegisterRequest','username为3～32位文字/数字/_/-；password 8～128；phone为含数字的6～30位电话；defaultLocation必填≤200。'],['CustomerLoginRequest','username和password必填，用户名最大32，密码最大128。'],['CustomerIntentRequest','name必填≤80，phone符合电话规则，meetingLocation必填≤200。'],['TradeResultRequest','success布尔必填；失败时disposition必须为VOID或REQUEUE。']])
 text+='\n\n## 响应数据类型\n\n'+table(['模型','字段'],[['CustomerView','id:number；username:string；phone:string；defaultLocation:string；createdAt:string。无密码字段。'],['CustomerSessionView','authenticated:boolean；customer:CustomerView或null。注册/登录成功自动建立客户会话。'],['CustomerOrderView','id/productId:number；productName:string；price:number；name/phone/meetingLocation:string；status:string；submittedAt:string；position:number或null；passcode:string或null；productStatus:string。'],['PageView<T>','items:T[]；total:number；page:number；size:number。'],['ProductView','id,name,price,description,images:string[],status,createdAt,stock,categoryId,categoryName,parentCategoryId。'],['CategoryView','id:number；name:string；parentId:number或null。'],['IntentView','id,name,phone,status,position,queueAt,meetingLocation,customerUsername；访客账号和地点可空。'],['WorkbenchView','product:ProductView；intents:IntentView[]；waitingCount:number；activeIntent:IntentView或null。'],['BuyerIntentView','name:string；phone:string；status:string；position:number或null；productStatus:string。'],['IssuedIntentView','passcode:string；position:number。']])
 text+='''\n\n## 客户购买与隔离规则

POST /api/customer/register成功201并自动登录，重复用户名忽略大小写返回409。POST /api/customer/login成功200，错误密码401。GET /api/customer/session在未登录时仍返回200以及authenticated=false、customer=null。POST /api/customer/logout返回204，只移除客户身份。

GET /api/customer/orders从会话读取本人编号，不接受其他客户编号。POST /api/customer/products/{id}/intents必须已登录，传入本次联系信息与地点，归属绑定当前会话客户；商品名称和价格在提交时保存快照。

GET /api/seller/customers支持q按用户名或电话搜索；GET /api/seller/customers/{id}/orders只允许卖家，购买记录中的passcode固定为null，不暴露客户操作凭据。

## 商品和队列状态

每件商品独立FIFO；成功交易减库存1，有库存恢复在售，售罄下架并结束该商品等待意向。失败VOID或REQUEUE不扣库存。交易冻结不能手动解除。最多9张相册图，描述使用清理后的富文本。

旧current接口需要唯一有效商品上下文，多商品时返回409；新增页面使用带商品编号路径。

## 密码方式

客户和卖家使用PBKDF2-HMAC-SHA256、600000次迭代、16字节随机盐、32字节结果。卖家改密使旧卖家会话失效；登录会话编号会轮换；身份与客户会话分开处理。

## 测试对应

首轮T01～T10与U01～U07覆盖基础功能；第二轮R201～R203覆盖三项改进；第三轮R301～R308覆盖A.1与A.2。各轮功能/接口目录、详细输入、状态断言和实际结果见tests目录。接口自动化通过不代表网页人工实例通过。
'''
 save(SOURCE/'backend/后端接口说明.md',text)
 return text
def test_documents(data):
 for bundle in data:
  r=bundle['round'];records=bundle['records'];api=[o for o in records if o['mode']=='automated'];func=[o for o in records if o['mode']=='manual'];passed=sum(o['result']=='通过' for o in api)
  overview=[]
  for case,title in bundle['cases'].items():
   selected=[o for o in records if o['caseId']==case];overview.append([case,title,sum(o['mode']=='manual' for o in selected),sum(o['mode']=='automated' for o in selected)])
  text=f'''# {bundle['title']}

版本 V{r}.3；整理日期 {DATE}。本版按实际验收轮次整理，采用一个总测试方案、总用例清单及逐项详细实例。

## 总测试方案

验收范围：{'首轮基础商品、访客意向、FIFO、交易和上传。' if r==1 else '第二轮联系信息修改、手动冻结与解冻、卖家改密。' if r==2 else '第三轮挑战需求A.1商品升级与A.2客户升级。'}

目录为 `{gen.ROUNDS[r].name}/功能测试或接口测试/总用例编号/第n个测试实例/`。每个总用例至少8个详细实例。共{len(records)}个实例：接口{len(api)}个，功能{len(func)}个。

环境：Java17、Maven3.6.3+、Node22+、Vue3与Element Plus；接口使用JUnit5、MockMvc及独立内存H2，每例清理测试数据，卖家初始admin/123456。页面测试应使用独立业务数据目录，避免修改日常data。

执行：`tests/run-iteration{r}-tests.cmd`；全部回归运行`tests/run-tests.cmd`。接口请求体从各实例请求JSON读取，上传从实例目录真实文件读取。动态编号、口令码和图片地址通过前一步响应capture变量传递，不能手填固定编号。

通过标准：请求状态码和响应字段逐项与断言一致；页面操作、提示、状态与预期一致并保存实际记录。退出标准：本轮接口无失败，页面实例全部人工执行后另行确认；当前不把未执行的页面用例标为通过。

## 总用例清单

{table(['编号','用例内容','功能实例数','接口实例数'],overview)}

## 执行结果

2026-10-23实际执行中接口实例已通过{passed}/{len(api)}，页面功能实例{len(func)}个待人工执行。首轮原有人工通过确认仅属于历史记录，不作为本次新整理实例的执行证据。

'''
  for kind in ('功能测试','接口测试'):
   text+=f'## {kind}详细用例\n\n'
   for case,title in bundle['cases'].items():
    cases=[o for o in records if o['caseId']==case and o['type']==kind]
    if not cases:continue
    text+=f'### {case} {title}\n\n'
    rows=[]
    for o in cases:
     # 详细操作和每一条原始请求另列在对应实例文件，正文列出真实输入与检查要点。
     info=f'数据：{o["dataSummary"]}<br>步骤：'+o['stepSummary'].replace('\n','<br>')
     rows.append([o['id']+' '+o['title'],info,o['expectedSummary'],o['result']+'<br>'+o['designMethod']+'<br>'+o['file']])
    text+=table(['实例与场景','详细数据和步骤','预期结果','结果 方法 文件'],rows)+'\n\n'
  text+='## 测试设计方法\n\n'+table(['方法','本项目中的具体使用'],gen.METHODS.items())+'\n\n极限输入通过边界值分析体现，不把本次字段边界测试称为性能压测。本版未执行并发压力、长时间稳定性和浏览器全量自动化，不以这些名称描述结果。\n'
  md=TESTS/(TITLES[r]+'.md');save(md,text);docx_from_md(md.with_suffix('.docx'),text,landscape=True)
  result={'iteration':r,'updatedAt':DATE,'scope':list(bundle['cases'].values()),'interface':{'total':len(api),'passed':passed,'failed':sum(o['result']=='失败' for o in api),'pending':sum(o['result']=='待执行' for o in api)},'functional':{'total':len(func),'passed':0,'pending':len(func)},'instances':[{'id':o['id'],'type':o['type'],'result':o['result'],'file':o['file']} for o in records]}
  gen.write(TESTS/('首轮测试结果.json' if r==1 else f'{ROUND_NAMES[r]}测试结果.json'),result)
  note=f'''# {ROUND_NAMES[r]}测试说明

版本V{r}.3；日期{DATE}。总用例和详细实例均按功能测试、接口测试分目录，每总用例至少8个实例。接口{len(api)}项，本次通过{passed}项；页面{len(func)}项待人工执行。

- [总方案与完整详细用例]({TITLES[r]}.md)
- [Word用例]({TITLES[r]}.docx)
- [Excel测试数据清单]({'测试用例数据清单.xlsx' if r==1 else ROUND_NAMES[r]+'测试用例数据清单.xlsx'})
- [实例目录]({gen.ROUNDS[r].name})

执行 `tests/run-iteration{r}-tests.cmd`。全部回归执行 `tests/run-tests.cmd`。接口测试只使用内存H2和backend/target测试上传目录。页面功能测试的前置数据、操作步骤和预期在每个功能实例目录中，实际结果需执行后填写。

{'首轮原有30个接口参数化实例保留，增加50个文件驱动实例；另有1个实例目录完整性检查。' if r==1 else '第二轮保留9项Java基础补充回归，并新增/整理24个文件驱动实例。' if r==2 else '第三轮64个接口实例覆盖A.1五项和A.2三项，包含边界、状态、权限隔离与历史快照。'}

三轮接口均针对当前生产源码回归，但首轮和第二轮只归属当时验收的功能；不会把第三轮功能回填为历史完成项。旧发布包和Docker镜像未重建。
'''
  save(TESTS/('测试说明.md' if r==1 else f'{ROUND_NAMES[r]}测试说明.md'),note)
 save(TESTS/'三轮测试总方案.md','# 在线购物系统三轮测试总方案\n\n整理日期'+DATE+'。每轮均以总方案统筹功能测试与接口测试，每个总用例每类至少8个详细实例。首轮U类为功能测试、T类为接口测试；第二、三轮同一总编号在两个分类目录下分别保存，功能实例编号加F以区分。\n\n'+table(['轮次','范围','接口实例','功能实例'],[[b['round'],'；'.join(b['cases'].values()),sum(o['mode']=='automated' for o in b['records']),sum(o['mode']=='manual' for o in b['records'])] for b in data])+'\n\n[首轮](在线购物系统首轮验收测试用例.md) · [第二轮](在线购物系统第二轮验收测试用例.md) · [第三轮](在线购物系统第三轮验收测试用例.md)\n\n'+table(['设计方法','原理与用途'],gen.METHODS.items()))
 return data

def subsystem():
 text='''# 在线购物系统子系统划分说明

版本V3.0；修订日期2026-10-24。沿用原五个子系统，第三轮扩展商品、分类、库存与客户能力。

## 1 子系统职责

'''+table(['编号','子系统','主要职责'],[['S1','买家交互','访客商品浏览搜索分类、口令码操作；客户注册登录、默认地点确认和本人购买历史。'],['S2','卖家工作台','多商品管理、分类管理、库存管理、商品级队列交易、改密及客户管理历史查询。'],['S3','商品与媒体','商品持久化、两级分类、库存与售罄状态、最多9张相册及富媒体清理和展示。'],['S4','购买意向与交易','按商品FIFO、冻结、姓名电话修改、撤销与交易；客户归属、下单地点和名称价格快照。'],['S5','数据与公共支持','H2、JPA、独立角色会话、PBKDF2、统一错误、文件上传、分页DTO及数据隔离。']])+'''

## 2 调用边界

S1/S2通过HTTP接口访问S3/S4；S3/S4使用S5的持久化、鉴权与统一错误。客户注册和登录由CustomerController及CustomerService提供；客户历史从服务端会话取得客户编号；卖家客户查询由SellerCustomerController提供并经过既有卖家拦截器。

## 3 前后端对应

'''+table(['功能','前端入口','后端服务'],[['商品浏览和意向','StorefrontView、catalogApi、customerApi','CatalogController、PublicController、StoreService'],['卖家商品与分类','SellerDashboardView、SellerCategories','SellerController、CatalogController、CatalogService'],['客户账号','CustomerAuthView、customerSession','CustomerController、CustomerService、CustomerRepository'],['客户历史','CustomerAccountView','CustomerController、CustomerService、IntentRepository'],['卖家客户管理','SellerCustomers','SellerCustomerController、CustomerService'],['密码和会话','SellerLoginView、账号管理、客户登录','SellerCredentialService、PasswordHasher、SellerAuthInterceptor']])+'''

## 4 数据与协作

products关联categories和purchase_intents，customer_accounts可关联多条purchase_intents；访客customer_id为空。客户历史价格名称在提交时保存快照。卖家和客户会话身份分离，任何DTO不返回密码哈希。

## 5 三轮验收归属

首轮为基础购物意向闭环；第二轮为联系信息修改、手动冻结解冻、卖家改密；第三轮为A.1商品升级与A.2客户升级。完整接口及返回模型见后端接口说明，详细功能和接口测试见tests目录。
'''
 save(ROOT/'subsystem/在线购物系统子系统划分说明.md',text);docx_from_md(ROOT/'subsystem/在线购物系统子系统划分说明.docx',text)
def other_documents(data):
 third='''# 第三轮功能说明

版本V3.0；日期2026-10-24。第三次迭代验收挑战需求A.1与A.2。前两轮功能继续回归；开发源码沿用project/mvp-v1.0目录。

## 功能与验收范围

'''+table(['编号','功能','前端入口'],[['R301','多商品发布和独立队列','商品管理、商品首页'],['R302','一级二级分类','分类管理、商品分类选择'],['R303','库存及售罄自动下架','商品库存、交易结果登记'],['R304','浏览、关键词搜索和分类分页','商品首页'],['R305','多图和富媒体描述','商品编辑、图片相册'],['R306','客户注册登录及默认交易地点','/customer/register、/customer/login'],['R307','客户购买历史及账号隔离','/customer/account'],['R308','卖家客户信息及历史查询','卖家后台客户管理']])+'''

## 运行

后端目录运行mvn spring-boot:run；前端目录运行npm ci和npm run dev；打开http://localhost:5173。新增客户注册、登录、我的购买入口和后台客户管理。

release和release-iteration2中的JAR、旧ZIP以及Docker镜像是历史版本；本次更新源码和文档测试，没有把旧包重标为第三轮发布。

## 测试成果物

'''+table(['成果物','位置'],[['总测试方案','tests/三轮测试总方案.md'],['第三轮完整测试','tests/在线购物系统第三轮验收测试用例.md及.docx'],['第三轮Excel','tests/第三轮测试用例数据清单.xlsx'],['第三轮实例','tests/第三轮验收测试数据/功能测试或接口测试/编号/第n个测试实例'],['接口说明','backend/后端接口说明.md'],['数据模型','backend/docs/数据模型图.md及PNG/SVG']])+'\n\n每个第三轮总用例各有8个接口实例和8个功能实例，共128项。接口结果见第三轮测试结果.json；页面人工结果另行填写。\n'
 save(SOURCE/'第三轮功能说明.md',third)
 save(SOURCE/'第二轮功能说明.md',"""# 第二轮功能说明

版本V2.3，整理日期2026-10-24。验收范围保持三项基础改进：R201等待买家修改联系信息；R202手动冻结与解冻；R203卖家改密和旧会话撤销。

挑战需求A.1与A.2统一由第三轮验收，见[第三轮功能说明](第三轮功能说明.md)。本次为第二轮补齐功能测试与接口测试各24个实例，每总用例每类8个。

本轮接口测试结果见tests/第二轮测试结果.json，页面人工用例待执行。先前2026-10-21的9个文件实例和9项补充回归通过仅作历史记录，当前完整数据以新目录为准。

开发启动：backend执行mvn spring-boot:run，frontend执行npm run dev。原release-iteration2中的app.jar属于第二轮历史集成包，不代表第三轮当前源码。
""")
 text='''# 在线购物系统

Vue3与Element Plus前端，Spring Boot与H2后端。当前第三次迭代完成挑战需求A.1商品管理升级与A.2客户升级，源码沿用project/mvp-v1.0。

## 三轮功能

'''+table(['轮次','验收内容'],[['首轮','访客浏览、购买意向与口令码、FIFO队列、交易结果及基础上传。'],['第二轮','等待联系信息修改、手动冻结解冻、卖家密码修改。'],['第三轮','多商品、两级分类、库存与售罄下架、搜索、多图富文本、客户注册登录和历史、卖家客户管理。']])+'''

## 开发启动

需要Java17+、Maven3.6.3+、Node22+。后端目录project/mvp-v1.0/backend执行mvn spring-boot:run；前端目录project/mvp-v1.0/frontend执行npm ci、npm run dev。打开http://localhost:5173。

'''+table(['入口','路径'],[['商品浏览与意向','/'],['客户注册','/customer/register'],['客户登录','/customer/login'],['我的购买历史','/customer/account'],['卖家登录','/seller/login'],['商品、分类、库存、队列、客户及账号管理','/seller/dashboard']])+'''

初始卖家admin/123456仅首次建库使用，改密后不会因重启恢复初始密码。日常H2数据库及上传目录在启动目录data中。

## 文档和成果物

'''+table(['成果物','链接'],[['需求规格','[完整需求](requirement/在线购物系统需求规格说明书.md)'],['第三轮需求补充','[第三轮需求](requirement/第三轮/第三轮需求补充说明.md)'],['第三轮功能说明','[功能说明](project/mvp-v1.0/第三轮功能说明.md)'],['子系统划分','[子系统说明](subsystem/在线购物系统子系统划分说明.md)'],['后端接口','[接口说明](project/mvp-v1.0/backend/后端接口说明.md)'],['数据模型','[字段与关系](project/mvp-v1.0/backend/docs/数据模型图.md)'],['三轮测试方案','[测试总方案](tests/三轮测试总方案.md)'],['首轮测试','[首轮用例](tests/在线购物系统首轮验收测试用例.md)'],['第二轮测试','[第二轮用例](tests/在线购物系统第二轮验收测试用例.md)'],['第三轮测试','[第三轮用例](tests/在线购物系统第三轮验收测试用例.md)']])+'''

## 测试运行

运行tests/run-tests.cmd执行全部接口回归和已有基础补充测试；也可运行tests/run-iteration1-tests.cmd、run-iteration2-tests.cmd、run-iteration3-tests.cmd分轮运行。

每轮实例目录下先分功能测试和接口测试，再按总用例编号及第n个测试实例划分。每总用例每类至少8个实例；Excel中分别展示总方案、功能测试、接口测试、原始请求数据，末行解释测试设计方法。

功能测试为网页人工验收，待执行状态单独保留；接口自动化使用独立内存H2与target测试上传目录，不操作日常data。旧发布包和Docker镜像未重建，运行本次功能请使用当前源码，不把历史JAR作为第三轮包。
'''
 save(ROOT/'README.md',text)
 save(SOURCE/'frontend/README.md','''# 在线购物系统前端

第三轮版本，Vue3、Element Plus、Vue Router、Axios。需要Node22及以上。

## 启动

frontend目录执行npm ci、npm run dev，访问http://localhost:5173；/api和/uploads代理至后端8080。构建命令npm run build。

## 页面与数据

StorefrontView负责多商品、搜索、分类、相册和购买意向；CustomerAuthView负责客户注册登录；CustomerAccountView负责本人历史；SellerDashboardView负责商品分类库存和队列，SellerCustomers负责客户查询与历史。customerSession维护客户身份，sellerApi和customerApi分别封装接口。

## 验收

三轮功能测试资料位于tests各轮数据目录的功能测试子目录，每例包含操作步骤、数据和预期。页面人工测试结果单独填写，不直接使用接口测试通过状态。
''')
 # 历史构建部署材料注明版本，避免将旧镜像误当第三轮。
 for p in (SOURCE/'release').glob('*.md'):
  s=p.read_text(encoding='utf-8-sig')
  if '历史发布资料' not in s:save(p,'历史发布资料：以下记录对应原发布版本；第三轮源码、客户功能与测试资料见[第三轮功能说明](../第三轮功能说明.md)。本次没有重建JAR、ZIP或Docker镜像。\n\n'+s)
 # 流程图可编辑源码同步到当前规则，不再包含唯一商品或视频已实现的说法。
 flow='''# 在线购物系统业务流程图

版本V3.0，2026-10-24。第三轮A.1与A.2，图中交易按指定商品独立处理。

## 商品与交易主流程

```mermaid
flowchart TD
 A[卖家发布多个商品] --> B[买家浏览搜索并选择商品]
 B --> C{客户是否登录}
 C -->|否| D[访客填写姓名电话]
 C -->|是| E[确认姓名电话及本次交易地点]
 D --> F[商品独立FIFO队列与口令码]
 E --> F
 E --> H[保存客户归属及商品快照]
 F --> G[卖家开始队首线下交易]
 G --> I{结果}
 I -->|失败| J[作废或重排 库存不变]
 I -->|成功| K[库存减1 当前意向成功]
 K --> L{库存为0}
 L -->|是| M[自动下架并结束其余等待]
 L -->|否| N[恢复在售 保留其余队列]
 J --> N
 H --> O[客户我的购买与卖家客户历史]
```

## 买家与客户历史

```mermaid
flowchart TD
 A[访客浏览搜索] --> B[注册或登录客户]
 B --> C[默认地点自动填入 可改单次地点]
 C --> D[提交指定商品意向]
 D --> E[本人购买历史]
 E --> F[等待时可凭码修改或撤销]
 E --> G[结束记录保留 无有效口令]
```

## 卖家客户管理

```mermaid
flowchart TD
 A[卖家登录] --> B[客户管理]
 B --> C[用户名或电话搜索]
 C --> D[客户详情及分页购买历史]
 D --> E[展示快照与交易地点 不展示密码或口令]
```

## 商品媒体

```mermaid
flowchart TD
 A[编辑商品名称价格库存分类] --> B[相册0至9张 JPG或PNG 单张不超过5MB]
 B --> C[富文本标题列表与本地图片]
 C --> D[服务端清理及校验]
 D --> E[按库存决定初始在售或下架]
```
'''
 save(ROOT/'process-and-state-diagrams/业务流程图.md',flow)
 states='''# 在线购物系统状态图

版本V3.0，2026-10-24。状态按商品隔离，客户历史独立于口令码有效性保留。

## 商品状态

```mermaid
stateDiagram-v2
 [*] --> ON_SALE: 有库存发布
 [*] --> DELISTED: 零库存发布
 ON_SALE --> MANUALLY_FROZEN: 手动冻结
 MANUALLY_FROZEN --> ON_SALE: 手动解冻
 ON_SALE --> IN_TRADE: 队首开始交易
 IN_TRADE --> ON_SALE: 成功有余量或失败
 IN_TRADE --> DELISTED: 成功售罄
 ON_SALE --> DELISTED: 主动下架或库存调为0
 DELISTED --> ON_SALE: 补货后重新上架
```

## 购买意向

```mermaid
stateDiagram-v2
 [*] --> WAITING: 提交
 WAITING --> WAITING: 修改联系信息或重排
 WAITING --> IN_TRADE: 队首交易
 IN_TRADE --> TRADE_SUCCESS: 成交
 IN_TRADE --> TRADE_FAILED: 失败作废
 IN_TRADE --> WAITING: 失败重排
 WAITING --> VOIDED: 卖家作废
 WAITING --> CANCELLED: 买家撤销
 WAITING --> TRADE_FAILED: 商品售罄
 WAITING --> PRODUCT_DELISTED: 主动下架
```

## 口令码

```mermaid
stateDiagram-v2
 [*] --> 有效: 提交意向
 有效 --> 有效: 等待冻结交易及失败重排
 有效 --> 失效: 当前意向结束或商品售罄下架
 失效 --> [*]: 不能凭原码操作 历史继续保留
```
'''
 save(ROOT/'process-and-state-diagrams/状态图.md',states)
def main():
 diagrams();requirement();api_and_model();subsystem();data=manifest();test_documents(data);other_documents(data)
 gen.write(TESTS/'测试资料索引.json',data)
 print('文档与索引同步完成；Word5份，三轮测试说明、需求与接口及模型已更新。')
if __name__=='__main__':main()
